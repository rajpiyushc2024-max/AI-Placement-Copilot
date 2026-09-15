import io
import pymupdf
from docx import Document


# ============================================================
# EXTRACT TEXT FROM PDF
# ============================================================

def extract_pdf_text(file):

    # Get complete file bytes
    file_bytes = file.getvalue()

    pdf = pymupdf.open(
        stream=file_bytes,
        filetype="pdf"
    )

    text = []

    for page in pdf:
        page_text = page.get_text()

        if page_text.strip():
            text.append(page_text.strip())

    pdf.close()

    return "\n".join(text).strip()


# ============================================================
# EXTRACT TEXT FROM DOCX
# ============================================================

def extract_docx_text(file):

    # Get complete file bytes
    file_bytes = file.getvalue()

    # Create a fresh in-memory file
    file_stream = io.BytesIO(file_bytes)

    document = Document(file_stream)

    text = []

    # --------------------------------------------------------
    # Normal paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            text.append(paragraph_text)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text.append(" | ".join(row_text))

    return "\n".join(text).strip()


# ============================================================
# MAIN TEXT EXTRACTION FUNCTION
# ============================================================

def extract_text(file):

    file_name = file.name.lower()

    if file_name.endswith(".pdf"):

        return extract_pdf_text(file)

    elif file_name.endswith(".docx"):

        return extract_docx_text(file)

    else:

        raise ValueError(
            "Only PDF and DOCX files are supported."
        )
